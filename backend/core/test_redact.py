"""Redacted invoice exports and PDF text-layer redaction."""
import io
import subprocess
import tempfile
from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from unittest import TestCase

from django.contrib.auth import get_user_model
from django.test import TestCase as DjangoTestCase
from reportlab.pdfgen import canvas

from core.invoicing import build_pdf, build_xlsx
from core.models import Account, Client, IssuerProfile
from core.redact import RedactError, redact_name, redact_pdf, secret_phrases


def _pdf_text(data: bytes) -> str:
    with tempfile.NamedTemporaryFile(suffix=".pdf") as fh:
        fh.write(data)
        fh.flush()
        return subprocess.check_output(["pdftotext", "-q", fh.name, "-"], text=True)


def _invoice():
    return SimpleNamespace(
        number="26092301",
        issuer_name="Oleksiy Zelenyuk",
        issuer_legal_form="Autónomo",
        issuer_city="Valencia",
        issuer_country="Spain",
        issuer_phone="34 605-817-513",
        issuer_email="o@zelen.uk",
        issuer_vat_number="ESZ0497876T",
        issuer_address="Carrer Picaio 9 p5",
        client_name="Quarrymare Limited",
        client_address="Dublin 15",
        client_tax_id="IE3388274LH",
        description="Senior Solutions Architect",
        quantity=Decimal("8"),
        unit="day",
        unit_price=Decimal("410"),
        net_amount=Decimal("3280"),
        vat_amount=Decimal("0"),
        vat_label="NP",
        total_amount=Decimal("3280"),
        currency="EUR",
        sale_date=date(2026, 9, 30),
        issue_date=date(2026, 9, 23),
        due_date=date(2026, 10, 7),
        iban="ES1915830001119099058403",
        bic="REVOESM2",
        correspondent_bic="",
        bank_address="Madrid",
        bank_name="Revolut",
        notes="",
    )


class InvoiceRedactTests(TestCase):
    def test_pdf_omits_bank_and_amounts_keeps_contacts(self):
        text = _pdf_text(build_pdf(_invoice(), redacted=True))
        self.assertIn("Oleksiy Zelenyuk", text)
        self.assertIn("Quarrymare", text)
        self.assertIn("o@zelen.uk", text)
        self.assertIn("ESZ0497876T", text)
        self.assertIn("Carrer Picaio 9 p5", text)
        self.assertIn("Dublin 15", text)
        self.assertNotIn("3 280,00", text)
        self.assertNotIn("410,00", text)
        self.assertNotIn("ES1915830001119099058403", text)
        self.assertNotIn("REVOESM2", text)
        self.assertIn("Redacted copy", text)

    def test_xlsx_omits_iban(self):
        from openpyxl import load_workbook

        data = build_xlsx(_invoice(), redacted=True)
        self.assertNotIn(b"ES1915830001119099058403", data)
        wb = load_workbook(io.BytesIO(data))
        values = [c.value for row in wb.active.iter_rows() for c in row if c.value]
        self.assertTrue(any("Redacted copy" in str(v) for v in values))
        self.assertNotIn(410.0, values)
        self.assertNotIn(3280.0, values)

    def test_full_pdf_still_has_iban(self):
        text = _pdf_text(build_pdf(_invoice()))
        self.assertIn("ES1915830001119099058403", text)


class PdfRedactTests(TestCase):
    def test_blacks_rate_and_words(self):
        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.drawString(72, 720, "Fee: 410 (Four Hundred and Ten Euro) per day")
        c.save()
        out = redact_pdf(
            buf.getvalue(),
            ["four hundred and ten euro"],
            amounts={Decimal("410")},
        )
        text = _pdf_text(out)
        self.assertNotIn("410", text)
        self.assertNotIn("Four Hundred", text)

    def test_blacks_iban(self):
        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.drawString(72, 720, "Please pay ES1915830001119099058403 today")
        c.save()
        out = redact_pdf(buf.getvalue(), ["ES1915830001119099058403"])
        self.assertTrue(out.startswith(b"%PDF"))
        self.assertNotIn("ES1915830001119099058403", _pdf_text(out))

    def test_rejects_when_nothing_matches(self):
        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.drawString(72, 720, "Nothing sensitive here")
        c.save()
        with self.assertRaises(RedactError):
            redact_pdf(buf.getvalue(), ["ES1915830001119099058403"])

    def test_redact_name(self):
        self.assertEqual(redact_name("contract.pdf"), "contract-redacted.pdf")


class SecretPhraseTests(DjangoTestCase):
    def test_collects_iban_and_amounts_not_contacts(self):
        user = get_user_model().objects.create_user("t", password="x")
        IssuerProfile.objects.create(
            owner=user,
            legal_name="Oleksiy Zelenyuk",
            vat_number="ESZ0497876T",
            phone="34605817513",
            email="o@zelen.uk",
            address="Carrer Picaio 9 p5, Puçol",
        )
        Account.objects.create(
            owner=user, name="Revolut", iban="ES1915830001119099058403"
        )
        Client.objects.create(
            owner=user, name="Q", default_unit_price=Decimal("410")
        )
        phrases = secret_phrases(user)
        blob = " ".join(phrases).lower()
        self.assertIn("es1915830001119099058403", blob)
        self.assertNotIn("esz0497876t", blob)
        self.assertNotIn("o@zelen.uk", blob)
        self.assertNotIn("picaio", blob)
        self.assertTrue(any("four hundred and ten" in p.lower() for p in phrases))
