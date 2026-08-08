"""Invoice helpers: working days, numbering, snapshots, XLSX/PDF export."""
from __future__ import annotations

import calendar
import io
import re
from datetime import date, timedelta
from decimal import Decimal
from functools import lru_cache
from pathlib import Path
from xml.sax.saxutils import escape

from django.core.files.base import ContentFile
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import Account, Client, Invoice, IssuerProfile

# Helvetica has no Polish glyphs (Ł/ł/ń → ■). Prefer DejaVu (Docker) / Arial (macOS).
_FONT_REGULAR_CANDIDATES = (
    Path(__file__).resolve().parent / "fonts" / "DejaVuSans.ttf",
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/dejavu/DejaVuSans.ttf"),
    Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
    Path("/Library/Fonts/Arial.ttf"),
)
_FONT_BOLD_CANDIDATES = (
    Path(__file__).resolve().parent / "fonts" / "DejaVuSans-Bold.ttf",
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    Path("/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"),
    Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    Path("/Library/Fonts/Arial Bold.ttf"),
)

REVERSE_CHARGE_NOTES = (
    "Invoice exempt from VAT pursuant to Directive 2006/112/EC and art. 25 "
    "of VAT Law 37/1992\n"
    "Operación con inversión del sujeto pasivo conforme al Art. 84 (Uno.2°) "
    "de la Ley 37/1992 de IVA\n"
    "Reverse charge*"
)

VAT_LABELS = {
    Client.VAT_REVERSE_CHARGE: "NP",
    Client.VAT_EXEMPT: "ZW",
    Client.VAT_STANDARD: "23%",
}


def working_days_in_month(year: int, month: int) -> int:
    """Count Mon–Fri days in a calendar month (ignores holidays)."""
    return sum(
        1
        for day in range(1, calendar.monthrange(year, month)[1] + 1)
        if date(year, month, day).weekday() < 5
    )


def advise_hours(year: int, month: int, hours_per_day: int = 8) -> dict:
    days = working_days_in_month(year, month)
    return {
        "year": year,
        "month": month,
        "working_days": days,
        "hours_per_day": hours_per_day,
        "suggested_hours": days * hours_per_day,
        "sale_date": date(year, month, calendar.monthrange(year, month)[1]).isoformat(),
    }


def last_day_of_month(year: int, month: int) -> date:
    return date(year, month, calendar.monthrange(year, month)[1])


def suggest_number(owner, issue_date: date) -> str:
    """YYMMDD## — next free seq for that issue date under the owner."""
    prefix = issue_date.strftime("%y%m%d")
    existing = (
        Invoice.objects.filter(owner=owner, number__startswith=prefix)
        .values_list("number", flat=True)
    )
    used = set()
    for num in existing:
        m = re.fullmatch(rf"{re.escape(prefix)}(\d{{2,}})", num)
        if m:
            used.add(int(m.group(1)))
    seq = 1
    while seq in used:
        seq += 1
    return f"{prefix}{seq:02d}"


def resolve_vat(client: Client) -> tuple[Decimal, str, str]:
    if client.vat_mode == Client.VAT_STANDARD:
        rate = client.default_vat_rate or Decimal("23")
        return rate, VAT_LABELS[Client.VAT_STANDARD], ""
    if client.vat_mode == Client.VAT_EXEMPT:
        return Decimal("0"), VAT_LABELS[Client.VAT_EXEMPT], ""
    return Decimal("0"), VAT_LABELS[Client.VAT_REVERSE_CHARGE], REVERSE_CHARGE_NOTES


def apply_snapshots(invoice: Invoice) -> None:
    """Copy issuer / client / account requisites onto the invoice row."""
    try:
        profile = invoice.owner.issuer_profile
    except IssuerProfile.DoesNotExist:
        profile = None
    if profile:
        for k, v in profile.invoice_snapshot().items():
            setattr(invoice, k, v)

    if invoice.client_id:
        for k, v in invoice.client.invoice_snapshot().items():
            setattr(invoice, k, v)
        rate, label, notes = resolve_vat(invoice.client)
        invoice.vat_rate = rate
        invoice.vat_label = label
        if not invoice.notes and notes:
            invoice.notes = notes

    if invoice.account_id:
        for k, v in invoice.account.invoice_snapshot().items():
            if k == "currency":
                # Keep invoice currency unless empty.
                if not invoice.currency:
                    invoice.currency = v
                continue
            setattr(invoice, k, v)


def default_invoice_account(owner) -> Account | None:
    return (
        Account.objects.filter(owner=owner, is_invoice_default=True).first()
        or Account.objects.filter(owner=owner, is_default=True).first()
        or Account.objects.filter(owner=owner).order_by("id").first()
    )


def _money(value) -> str:
    return f"{Decimal(value):,.2f}".replace(",", "X").replace(".", ",").replace("X", " ")


def _date(d: date | None) -> str:
    return d.strftime("%d.%m.%Y") if d else ""


def build_xlsx(invoice: Invoice) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"

    bold = Font(bold=True, size=12)
    title = Font(bold=True, size=16)
    small = Font(size=9)
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    ws["D3"] = invoice.issuer_name
    ws["D3"].font = title
    ws["D5"] = invoice.issuer_legal_form
    ws["G5"] = ", ".join(
        p for p in [invoice.issuer_city, invoice.issuer_country] if p
    )

    ws["B7"] = f"Invoice #{invoice.number}"
    ws["B7"].font = bold
    ws["D7"] = "Contact/Contacto"
    ws["E7"] = "Customer/Cliente"
    ws["D7"].font = bold
    ws["E7"].font = bold

    ws["D8"] = invoice.issuer_phone
    ws["E8"] = invoice.client_name
    ws["D9"] = invoice.issuer_email
    ws["D10"] = f"VAT# {invoice.issuer_vat_number}" if invoice.issuer_vat_number else ""
    ws["D11"] = invoice.issuer_address
    client_line = invoice.client_address
    if invoice.client_tax_id:
        client_line = f"{client_line} NIP: {invoice.client_tax_id}".strip()
    ws["E11"] = client_line

    ws["B13"] = "Salesperson / Vendedor"
    ws["D13"] = "Description"
    ws["E13"] = "Cantidad"
    ws["F13"] = "Precio/unidad"
    ws["G13"] = "Total de línea"
    for col in ("B", "D", "E", "F", "G"):
        ws[f"{col}13"].font = bold
        ws[f"{col}13"].border = thin

    ws["B14"] = invoice.issuer_name
    ws["D14"] = invoice.description
    ws["E14"] = float(invoice.quantity)
    ws["F14"] = float(invoice.unit_price)
    ws["G14"] = float(invoice.net_amount)
    ws["F14"].number_format = "0.00"
    ws["G14"].number_format = '#,##0.00'

    ws["B16"] = "Date of Sale / Fecha de vendido"
    ws["B17"] = invoice.sale_date
    ws["B17"].number_format = "DD.MM.YYYY"
    ws["B19"] = "Date of Issue / Fecha de emisión"
    ws["B20"] = invoice.issue_date
    ws["B20"].number_format = "DD.MM.YYYY"
    ws["B22"] = "Due Date / Fecha limite"
    ws["B23"] = invoice.due_date
    ws["B23"].number_format = "DD.MM.YYYY"

    ws["D22"] = f"IBAN: {invoice.iban}" if invoice.iban else ""
    ws["D23"] = f"BIC: {invoice.bic}" if invoice.bic else ""
    if invoice.correspondent_bic:
        ws["D24"] = f"Correspondent BIC: {invoice.correspondent_bic}"
    ws["E24"] = invoice.bank_address or invoice.bank_name

    ws["E21"] = "Subtotal"
    ws["F22"] = f"VAT ({invoice.vat_label})" if invoice.vat_label else "VAT"
    ws["G22"] = "Total"
    for cell in ("E21", "F22", "G22"):
        ws[cell].font = bold
    ws["E23"] = float(invoice.net_amount)
    ws["F23"] = float(invoice.vat_amount)
    ws["G23"] = float(invoice.total_amount)
    for cell in ("E23", "F23", "G23"):
        ws[cell].number_format = '#,##0.00'
        ws[cell].font = bold

    if invoice.notes:
        row = 25
        for line in invoice.notes.splitlines():
            ws[f"C{row}"] = line
            ws[f"C{row}"].font = small
            row += 1

    ws["H27"] = "THANK YOU / GRACIAS"
    ws["H27"].font = bold
    ws["H27"].alignment = Alignment(horizontal="right")

    for col, width in {
        "A": 3, "B": 28, "C": 12, "D": 36, "E": 18, "F": 14, "G": 14, "H": 18
    }.items():
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@lru_cache(maxsize=1)
def _pdf_fonts() -> tuple[str, str]:
    """Register a Unicode TTF pair; fall back to Helvetica only if nothing found."""
    regular = next((p for p in _FONT_REGULAR_CANDIDATES if p.is_file()), None)
    bold = next((p for p in _FONT_BOLD_CANDIDATES if p.is_file()), None)
    if regular is None:
        return "Helvetica", "Helvetica-Bold"
    pdfmetrics.registerFont(TTFont("InvoiceSans", str(regular)))
    if bold is not None:
        pdfmetrics.registerFont(TTFont("InvoiceSans-Bold", str(bold)))
        return "InvoiceSans", "InvoiceSans-Bold"
    return "InvoiceSans", "InvoiceSans"


def _p(text: str, style) -> Paragraph:
    return Paragraph(escape(text or "").replace("\n", "<br/>"), style)


def build_pdf(invoice: Invoice) -> bytes:
    font, font_bold = _pdf_fonts()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle(
        "h1",
        parent=styles["Heading1"],
        fontName=font_bold,
        fontSize=16,
        spaceAfter=4,
    )
    h2 = ParagraphStyle(
        "h2",
        parent=styles["Heading2"],
        fontName=font_bold,
        fontSize=11,
        spaceBefore=10,
        spaceAfter=4,
    )
    body = ParagraphStyle(
        "body", parent=styles["Normal"], fontName=font, fontSize=9, leading=12
    )
    body_bold = ParagraphStyle("body_bold", parent=body, fontName=font_bold)
    header = ParagraphStyle(
        "hdr", parent=body_bold, textColor=colors.white
    )
    small = ParagraphStyle(
        "small",
        parent=styles["Normal"],
        fontName=font,
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#444444"),
    )

    story = []
    story.append(_p(invoice.issuer_name or "Invoice", h1))
    meta = " · ".join(
        p
        for p in [
            invoice.issuer_legal_form,
            invoice.issuer_city,
            invoice.issuer_country,
        ]
        if p
    )
    if meta:
        story.append(_p(meta, body))
    story.append(_p(f"Invoice #{invoice.number}", h2))

    left = [
        _p("Contact", body_bold),
        _p(invoice.issuer_phone or "—", body),
        _p(invoice.issuer_email or "—", body),
        _p(
            f"VAT# {invoice.issuer_vat_number}" if invoice.issuer_vat_number else "—",
            body,
        ),
        _p(invoice.issuer_address or "", body),
    ]
    right = [
        _p("Customer", body_bold),
        _p(invoice.client_name or "—", body),
        _p(invoice.client_address or "", body),
        _p(f"NIP: {invoice.client_tax_id}" if invoice.client_tax_id else "", body),
    ]
    story.append(
        Table(
            [[left, right]],
            colWidths=[90 * mm, 80 * mm],
            hAlign="LEFT",
        )
    )
    story.append(Spacer(1, 8 * mm))

    qty = (
        str(int(invoice.quantity))
        if Decimal(invoice.quantity) == int(invoice.quantity)
        else str(invoice.quantity)
    )
    data = [
        [
            _p("Description", header),
            _p("Qty", header),
            _p("Unit price", header),
            _p("Line total", header),
        ],
        [
            _p(invoice.description or "", body),
            _p(qty, body),
            _p(_money(invoice.unit_price), body),
            _p(_money(invoice.net_amount), body),
        ],
    ]
    table = Table(data, colWidths=[90 * mm, 25 * mm, 30 * mm, 30 * mm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 6 * mm))

    dates = [
        [_p("Date of sale", body), _p(_date(invoice.sale_date), body)],
        [_p("Date of issue", body), _p(_date(invoice.issue_date), body)],
        [_p("Due date", body), _p(_date(invoice.due_date), body)],
        [
            _p("Subtotal", body),
            _p(f"{_money(invoice.net_amount)} {invoice.currency}", body),
        ],
        [
            _p(f"VAT ({invoice.vat_label})", body),
            _p(f"{_money(invoice.vat_amount)} {invoice.currency}", body),
        ],
        [
            _p("Total", body_bold),
            _p(f"{_money(invoice.total_amount)} {invoice.currency}", body_bold),
        ],
    ]
    bank = [
        f"IBAN: {invoice.iban}" if invoice.iban else "",
        f"BIC: {invoice.bic}" if invoice.bic else "",
        (
            f"Correspondent BIC: {invoice.correspondent_bic}"
            if invoice.correspondent_bic
            else ""
        ),
        invoice.bank_address or invoice.bank_name or "",
    ]
    bank_para = _p("\n".join(p for p in bank if p) or "—", body)
    dates_table = Table(dates, colWidths=[40 * mm, 45 * mm])
    dates_table.setStyle(
        TableStyle(
            [
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    story.append(Table([[dates_table, bank_para]], colWidths=[90 * mm, 80 * mm]))

    if invoice.notes:
        story.append(Spacer(1, 6 * mm))
        for line in invoice.notes.splitlines():
            story.append(_p(line, small))

    story.append(Spacer(1, 10 * mm))
    story.append(_p("THANK YOU / GRACIAS", h2))

    doc.build(story)
    return buf.getvalue()


def attach_generated_files(invoice: Invoice) -> None:
    stem = re.sub(r"[^\w.\-]+", "_", f"{invoice.number}")[:80] or "invoice"
    invoice.xlsx_file.save(f"{stem}.xlsx", ContentFile(build_xlsx(invoice)), save=False)
    invoice.pdf_file.save(f"{stem}.pdf", ContentFile(build_pdf(invoice)), save=False)


def issue_invoice(invoice: Invoice) -> Invoice:
    if invoice.status == Invoice.STATUS_ISSUED:
        raise ValueError("Invoice is already issued.")
    apply_snapshots(invoice)
    invoice.recalculate_amounts()
    if not invoice.number:
        invoice.number = suggest_number(invoice.owner, invoice.issue_date)
    attach_generated_files(invoice)
    invoice.status = Invoice.STATUS_ISSUED
    invoice.issued_at = timezone.now()
    invoice.save()
    return invoice


def parse_invoice_xlsx(path: str | Path) -> dict:
    """Best-effort parse of the Netguru-style invoice workbook."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb.active
    raw_number = str(ws["B7"].value or "")
    number = raw_number.replace("Invoice #", "").strip()
    qty = ws["E14"].value
    price = ws["F14"].value
    net = ws["G14"].value or ws["E23"].value
    sale = ws["B17"].value
    issue = ws["B20"].value
    due = ws["B23"].value

    def as_date(v):
        if hasattr(v, "date"):
            return v.date() if hasattr(v, "hour") else v
        return None

    sale_d, issue_d, due_d = as_date(sale), as_date(issue), as_date(due)
    # Prefer filename DD-MM-YYYY when Excel used TODAY() formulas.
    name = Path(path).name
    m = re.match(r"(\d{2})-(\d{2})-(\d{4})", name)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        file_issue = date(y, mo, d)
        # If cell issue looks like a TODAY()-relative ghost, trust filename.
        if issue_d is None or abs((issue_d - file_issue).days) > 10:
            issue_d = file_issue
            due_d = issue_d + timedelta(days=14)
            # Sale = last day of previous month
            first = date(issue_d.year, issue_d.month, 1)
            sale_d = first - timedelta(days=1)

    if sale_d is None and issue_d is not None:
        first = date(issue_d.year, issue_d.month, 1)
        sale_d = first - timedelta(days=1)
    if due_d is None and issue_d is not None:
        due_d = issue_d + timedelta(days=14)

    return {
        "number": number or suggest_number_from_dates(issue_d),
        "description": ws["D14"].value or "",
        "quantity": Decimal(str(qty or 0)),
        "unit_price": Decimal(str(price or 0)),
        "net_amount": Decimal(str(net or 0)),
        "vat_amount": Decimal(str(ws["F23"].value or 0)),
        "total_amount": Decimal(str(ws["G23"].value or net or 0)),
        "sale_date": sale_d,
        "issue_date": issue_d,
        "due_date": due_d,
        "iban": _strip_prefix(ws["D22"].value, "IBAN:"),
        "bic": _strip_prefix(ws["D23"].value, "BIC:"),
        "correspondent_bic": _strip_prefix(ws["D24"].value, "Correspondent BIC:"),
        "bank_address": ws["E24"].value or "",
        "client_name": ws["E8"].value or "",
        "client_address": ws["E11"].value or "",
        "issuer_name": ws["D3"].value or "",
        "issuer_phone": ws["D8"].value or "",
        "issuer_email": ws["D9"].value or "",
        "issuer_vat_number": _strip_prefix(ws["D10"].value, "VAT#"),
        "issuer_address": ws["D11"].value or "",
        "issuer_legal_form": ws["D5"].value or "",
    }


def suggest_number_from_dates(issue_date: date | None) -> str:
    if not issue_date:
        return date.today().strftime("%y%m%d") + "01"
    return issue_date.strftime("%y%m%d") + "01"


def _strip_prefix(value, prefix: str) -> str:
    text = str(value or "").strip()
    if text.upper().startswith(prefix.upper()):
        return text[len(prefix) :].strip()
    return text
