"""Export folder transactions to an .xlsx file with per-sheet totals."""
from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADERS = ["Fecha operación", "Fecha valor", "Concepto", "Contraparte", "Importe", "Saldo", "Moneda"]


def _sanitize_sheet_title(name: str) -> str:
    invalid = set(r"[]:*?/\\")
    cleaned = "".join(c for c in name if c not in invalid).strip()
    return (cleaned or "Folder")[:31]


def _unique_title(title: str, used: set[str]) -> str:
    """Excel sheet titles must be unique and <= 31 chars."""
    base = _sanitize_sheet_title(title)
    if base not in used:
        used.add(base)
        return base
    for i in range(2, 1000):
        suffix = f" ({i})"
        candidate = base[: 31 - len(suffix)] + suffix
        if candidate not in used:
            used.add(candidate)
            return candidate
    used.add(base)
    return base


def _write_sheet(ws, transactions, totals: str) -> None:
    ws.append(HEADERS)
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    income = Decimal("0")
    expense = Decimal("0")
    transactions = list(transactions)
    currency = transactions[0].currency if transactions else "EUR"

    for tx in transactions:
        ws.append(
            [
                tx.operation_date,
                tx.value_date,
                tx.concept,
                tx.counterparty,
                float(tx.amount),
                float(tx.balance) if tx.balance is not None else None,
                tx.currency,
            ]
        )
        if tx.amount >= 0:
            income += tx.amount
        else:
            expense += tx.amount
        row_idx = ws.max_row
        ws.cell(row=row_idx, column=1).number_format = "DD/MM/YYYY"
        ws.cell(row=row_idx, column=2).number_format = "DD/MM/YYYY"
        ws.cell(row=row_idx, column=5).number_format = "#,##0.00"
        ws.cell(row=row_idx, column=6).number_format = "#,##0.00"

    ws.append([])
    if totals in ("income", "both"):
        _append_total(ws, "Total ingresos", income, currency)
    if totals in ("expense", "both"):
        _append_total(ws, "Total gastos", expense, currency)
    if totals == "both":
        _append_total(ws, "Neto", income + expense, currency)

    _autofit(ws)


def build_workbook(sheets, totals: str = "both") -> io.BytesIO:
    """Build an .xlsx from `sheets`, a list of (title, transactions) tuples.

    Each tuple becomes its own worksheet with its own totals. `totals` is one
    of "income", "expense", "both".
    """
    sheets = list(sheets)
    wb = Workbook()
    used: set[str] = set()

    if not sheets:
        _write_sheet(wb.active, [], totals)
        return _save(wb)

    for i, (title, transactions) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = _unique_title(title, used)
        _write_sheet(ws, transactions, totals)

    return _save(wb)


def build_folder_workbook(title, transactions, totals: str = "both") -> io.BytesIO:
    """Single-sheet workbook for one list of transactions (backwards compat)."""
    return build_workbook([(title, transactions)], totals=totals)


def _save(wb) -> io.BytesIO:
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _append_total(ws, label: str, value: Decimal, currency: str):
    ws.append(["", "", label, "", float(value), None, currency])
    row_idx = ws.max_row
    ws.cell(row=row_idx, column=3).font = TOTAL_FONT
    amount_cell = ws.cell(row=row_idx, column=5)
    amount_cell.font = TOTAL_FONT
    amount_cell.number_format = "#,##0.00"


def _autofit(ws):
    widths = [16, 16, 50, 28, 14, 14, 8]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
